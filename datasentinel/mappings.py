from __future__ import annotations

import csv
import json
import operator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from types import MappingProxyType
from typing import Any, Callable, Mapping


@dataclass(frozen=True)
class ColumnMapping:
    source_columns: tuple[str, ...]
    target_column: str
    operation: str = "copy"
    options: Mapping[str, Any] = field(default_factory=dict)
    target_type: str | None = None
    target_length: int | None = None
    nullable: bool = True
    default: Any = None
    has_default: bool = False

    def __post_init__(self) -> None:
        object.__setattr__(self, "options", MappingProxyType(dict(self.options)))

    def apply(self, source_record: Mapping[str, Any]) -> Any:
        missing = [column for column in self.source_columns if column not in source_record]
        if missing:
            raise ValueError(
                f"Mapping for target '{self.target_column}' requires missing source column(s): {missing}."
            )
        values = [source_record[column] for column in self.source_columns]
        result = _apply_operation(self.operation, values, self.options)
        if result is None and self.has_default:
            result = self.default
        if result is None and not self.nullable:
            raise ValueError(
                f"Mapping for non-nullable target '{self.target_column}' produced null."
            )
        if result is not None and self.target_type is not None:
            result = _cast_value(result, self.target_type, self.target_column)
        if (
            result is not None
            and self.target_length is not None
            and isinstance(result, str)
            and len(result) > self.target_length
        ):
            raise ValueError(
                f"Value for target '{self.target_column}' exceeds its configured length "
                f"of {self.target_length}."
            )
        return result


@dataclass(frozen=True)
class MappingDocument:
    name: str
    mappings: tuple[ColumnMapping, ...]
    source_columns: tuple[str, ...] = ()
    target_columns: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        self.validate()

    def validate(self) -> None:
        if not self.name.strip():
            raise ValueError("Mapping document name must not be empty.")
        if not self.mappings:
            raise ValueError("Mapping document must define at least one column mapping.")

        target_names: set[str] = set()
        for item in self.mappings:
            if not item.target_column.strip():
                raise ValueError("Each column mapping must define a target_column.")
            if item.target_column in target_names:
                raise ValueError(f"Target column '{item.target_column}' is mapped more than once.")
            target_names.add(item.target_column)
            if not item.source_columns:
                raise ValueError(f"Mapping for target '{item.target_column}' requires source_columns.")
            if any(not column.strip() for column in item.source_columns):
                raise ValueError(f"Mapping for target '{item.target_column}' has an empty source column.")
            if item.operation not in _SUPPORTED_OPERATIONS:
                raise ValueError(
                    f"Unsupported transformation '{item.operation}' for target '{item.target_column}'. "
                    f"Supported transformations: {', '.join(sorted(_SUPPORTED_OPERATIONS))}."
                )
            _validate_operation(item)
            if item.target_type is not None and _normalize_target_type(item.target_type) not in _SUPPORTED_TYPES:
                raise ValueError(
                    f"Unsupported target_type '{item.target_type}' for target '{item.target_column}'."
                )
            if item.target_length is not None and item.target_length <= 0:
                raise ValueError(f"target_length for '{item.target_column}' must be positive.")
        declared_sources = set(self.source_columns)
        declared_targets = set(self.target_columns)
        if len(declared_sources) != len(self.source_columns):
            raise ValueError("source_columns contains duplicate column names.")
        if len(declared_targets) != len(self.target_columns):
            raise ValueError("target_columns contains duplicate column names.")

        referenced_sources = {
            source_column
            for item in self.mappings
            for source_column in item.source_columns
        }
        undeclared_sources = referenced_sources - declared_sources
        if self.source_columns and undeclared_sources:
            raise ValueError(
                f"Mappings reference undeclared source column(s): {sorted(undeclared_sources)}."
            )
        undeclared_targets = target_names - declared_targets
        if self.target_columns and undeclared_targets:
            raise ValueError(
                f"Mappings reference undeclared target column(s): {sorted(undeclared_targets)}."
            )

    def validate_source_schema(self, actual_columns: list[str]) -> dict[str, list[str]]:
        expected = self.source_columns or list(dict.fromkeys(
            source for item in self.mappings for source in item.source_columns
        ))
        return _compare_columns(expected, actual_columns)

    def validate_target_schema(self, actual_columns: list[str]) -> dict[str, list[str]]:
        expected = self.target_columns or [item.target_column for item in self.mappings]
        return _compare_columns(expected, actual_columns)

    def build_transformation_rules(self) -> dict[str, Callable[[Mapping[str, Any]], Any]]:
        return {
            item.target_column: item.apply
            for item in self.mappings
        }

    def transform_record(self, source_record: Mapping[str, Any]) -> dict[str, Any]:
        rules = self.build_transformation_rules()
        return {
            target_column: rule(source_record)
            for target_column, rule in rules.items()
        }

    def summary(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "source_columns": list(self.source_columns or tuple(dict.fromkeys(
                source for item in self.mappings for source in item.source_columns
            ))),
            "target_columns": list(self.target_columns or tuple(
                item.target_column for item in self.mappings
            )),
            "transformation_rules": [
                {
                    "target_column": item.target_column,
                    "source_columns": list(item.source_columns),
                    "operation": item.operation,
                    "target_type": item.target_type,
                    "target_length": item.target_length,
                    "nullable": item.nullable,
                }
                for item in self.mappings
            ],
        }


_SUPPORTED_OPERATIONS = {
    "copy",
    "trim",
    "upper",
    "lower",
    "concat",
    "coalesce",
    "add",
    "subtract",
    "multiply",
    "divide",
    "replace",
    "substring",
    "round",
    "if_else",
}
_SUPPORTED_TYPES = {"string", "integer", "number", "decimal", "boolean", "date"}
_ARITHMETIC_OPERATIONS: dict[str, Callable[[Any, Any], Any]] = {
    "add": operator.add,
    "subtract": operator.sub,
    "multiply": operator.mul,
    "divide": operator.truediv,
}
_OPERATION_OPTIONS = {
    "copy": set(),
    "trim": set(),
    "upper": set(),
    "lower": set(),
    "concat": {"separator"},
    "coalesce": set(),
    "add": set(),
    "subtract": set(),
    "multiply": set(),
    "divide": set(),
    "replace": {"old", "new"},
    "substring": {"start", "length"},
    "round": {"digits"},
    "if_else": {"equals", "then", "else"},
}


def _validate_operation(item: ColumnMapping) -> None:
    operation = item.operation
    allowed_options = _OPERATION_OPTIONS[operation]
    unsupported_options = set(item.options) - allowed_options
    if unsupported_options:
        raise ValueError(
            f"Unsupported option(s) for transformation '{operation}' on target "
            f"'{item.target_column}': {sorted(unsupported_options)}."
        )

    source_count = len(item.source_columns)
    exact_counts = {
        "copy": 1,
        "trim": 1,
        "upper": 1,
        "lower": 1,
        "replace": 1,
        "substring": 1,
        "round": 1,
        "if_else": 1,
        "subtract": 2,
        "divide": 2,
    }
    if operation in exact_counts and source_count != exact_counts[operation]:
        raise ValueError(
            f"Transformation '{operation}' for target '{item.target_column}' requires "
            f"{exact_counts[operation]} source column(s)."
        )
    if operation in {"concat", "coalesce", "add", "multiply"} and source_count < 1:
        raise ValueError(
            f"Transformation '{operation}' for target '{item.target_column}' requires source columns."
        )
    if operation in {"add", "multiply"} and source_count < 2:
        raise ValueError(
            f"Transformation '{operation}' for target '{item.target_column}' requires at least two source columns."
        )
    if operation == "concat" and not isinstance(item.options.get("separator", ""), str):
        raise ValueError(f"Transformation 'concat' for target '{item.target_column}' requires a string separator.")
    if operation == "replace":
        if not isinstance(item.options.get("old"), str) or not isinstance(item.options.get("new", ""), str):
            raise ValueError(f"Transformation 'replace' for target '{item.target_column}' requires string old/new options.")
    if operation == "substring":
        start = item.options.get("start", 0)
        length = item.options.get("length")
        if isinstance(start, bool) or not isinstance(start, int):
            raise ValueError(f"Transformation 'substring' for target '{item.target_column}' requires an integer start.")
        if length is not None and (isinstance(length, bool) or not isinstance(length, int) or length < 0):
            raise ValueError(f"Transformation 'substring' for target '{item.target_column}' requires a non-negative integer length.")
    if operation == "round":
        digits = item.options.get("digits", 0)
        if isinstance(digits, bool) or not isinstance(digits, int):
            raise ValueError(f"Transformation 'round' for target '{item.target_column}' requires an integer digits option.")
    if operation == "if_else" and not {"equals", "then", "else"}.issubset(item.options):
        raise ValueError(
            f"Transformation 'if_else' for target '{item.target_column}' requires equals, then, and else options."
        )


def load_mapping_document(path: str | Path) -> MappingDocument:
    mapping_path = Path(path)
    suffix = mapping_path.suffix.lower()
    if suffix == ".json":
        raw = json.loads(mapping_path.read_text(encoding="utf-8"))
    elif suffix in {".yaml", ".yml"}:
        try:
            import yaml
        except ImportError as error:
            raise RuntimeError(
                "YAML mapping files require PyYAML. Install with: pip install data-sentinel[mappings]"
            ) from error
        try:
            raw = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
        except yaml.YAMLError as error:
            raise ValueError(f"Invalid YAML in mapping document '{mapping_path}'.") from error
    elif suffix == ".csv":
        with mapping_path.open("r", encoding="utf-8-sig", newline="") as stream:
            raw = {"name": mapping_path.stem, "mappings": list(csv.DictReader(stream))}
    elif suffix == ".xlsx":
        raw = _read_excel_mapping(mapping_path)
    else:
        raise ValueError(
            f"Unsupported mapping document format '{suffix}'. Supported formats: "
            ".json, .yaml, .yml, .csv, .xlsx."
        )
    return parse_mapping_document(raw, default_name=mapping_path.stem)


def parse_mapping_document(raw: Any, default_name: str = "uploaded-mapping") -> MappingDocument:
    if not isinstance(raw, dict):
        raise ValueError("Mapping document must contain an object at its root.")

    raw_mappings = raw.get("mappings", raw.get("column_mappings"))
    if not isinstance(raw_mappings, list):
        raise ValueError("Mapping document must contain a 'mappings' array.")

    mappings = tuple(_parse_column_mapping(entry, index) for index, entry in enumerate(raw_mappings))
    return MappingDocument(
        name=_nonempty_string(raw.get("name", default_name), "name"),
        mappings=mappings,
        source_columns=_parse_column_list(raw.get("source_columns", []), "source_columns"),
        target_columns=_parse_column_list(raw.get("target_columns", []), "target_columns"),
    )


def _parse_column_mapping(entry: Any, index: int) -> ColumnMapping:
    if not isinstance(entry, dict):
        raise ValueError(f"Mapping at index {index} must be an object.")

    entry = {
        key: value
        for key, value in _normalize_mapping_keys(entry).items()
        if not (isinstance(value, str) and not value.strip())
    }
    target_column = _nonempty_string(entry.get("target_column"), f"mappings[{index}].target_column")
    source_columns_value = entry.get("source_columns", entry.get("source_column"))
    if isinstance(source_columns_value, str):
        if source_columns_value.strip().startswith("["):
            try:
                source_columns_value = json.loads(source_columns_value)
            except json.JSONDecodeError as error:
                raise ValueError(
                    f"source_columns for target '{target_column}' must be a JSON array."
                ) from error
        else:
            source_columns_value = _nonempty_string(
                source_columns_value,
                f"mappings[{index}].source_column",
            )
    if isinstance(source_columns_value, str):
        source_columns = (source_columns_value,)
    elif isinstance(source_columns_value, list):
        source_columns = tuple(
            _nonempty_string(value, f"mappings[{index}].source_columns")
            for value in source_columns_value
        )
    else:
        raise ValueError(
            f"Mapping for target '{target_column}' must define source_column or source_columns."
        )

    transform = entry.get("transform", entry.get("transformation", "copy"))
    options: dict[str, Any] = {}
    if isinstance(transform, str):
        transform = transform.strip()
        if transform.startswith("{"):
            try:
                transform = json.loads(transform)
            except json.JSONDecodeError as error:
                raise ValueError(
                    f"Transformation for target '{target_column}' must be valid JSON."
                ) from error
    if isinstance(transform, str):
        operation = transform.lower()
    elif isinstance(transform, dict):
        operation = _nonempty_string(
            transform.get("op"),
            f"transform.op for target '{target_column}'",
        ).lower()
        options = {key: value for key, value in transform.items() if key != "op"}
    else:
        raise ValueError(f"Transformation for target '{target_column}' must be a string or object.")

    target_type = entry.get("target_type", options.pop("type", None))
    if target_type is not None:
        target_type = _normalize_target_type(
            _nonempty_string(target_type, f"target_type for '{target_column}'").lower()
        )
    target_length = entry.get("target_length", entry.get("max_length"))
    if target_length is not None:
        try:
            target_length = int(target_length)
        except (TypeError, ValueError) as error:
            raise ValueError(f"target_length for '{target_column}' must be a positive integer.") from error
    nullable_value = entry.get("nullable", True)
    if isinstance(nullable_value, str):
        normalized = nullable_value.strip().lower()
        if normalized not in {"true", "false"}:
            raise ValueError(f"nullable for target '{target_column}' must be true or false.")
        nullable_value = normalized == "true"
    if not isinstance(nullable_value, bool):
        raise ValueError(f"nullable for target '{target_column}' must be a boolean.")

    return ColumnMapping(
        source_columns=source_columns,
        target_column=target_column,
        operation=operation,
        options=options,
        target_type=target_type,
        target_length=target_length,
        nullable=nullable_value,
        default=entry.get("default"),
        has_default="default" in entry,
    )


def _normalize_mapping_keys(entry: Mapping[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in entry.items():
        if not isinstance(key, str):
            raise ValueError("Mapping field names must be strings.")
        normalized_key = "_".join(key.strip().lower().replace("-", " ").split())
        aliases = {
            "source": "source_column",
            "source_field": "source_column",
            "source_columns": "source_columns",
            "target": "target_column",
            "target_field": "target_column",
            "transformation_rule": "transform",
            "target_data_type": "target_type",
            "data_type": "target_type",
            "length": "target_length",
            "max_length": "target_length",
            "is_nullable": "nullable",
        }
        normalized[aliases.get(normalized_key, normalized_key)] = value
    return normalized


def _read_excel_mapping(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel mapping files require openpyxl. Install with: pip install data-sentinel[mappings]"
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError(f"Excel mapping document '{path}' is empty.")
        normalized_headers = [
            str(value).strip() if value is not None else ""
            for value in headers
        ]
        if any(not header for header in normalized_headers):
            raise ValueError("Excel mapping header row contains an empty column name.")
        entries = [
            dict(zip(normalized_headers, row))
            for row in rows
            if any(value is not None for value in row)
        ]
        return {"name": path.stem, "mappings": entries}
    finally:
        workbook.close()


def _parse_column_list(value: Any, field_name: str) -> tuple[str, ...]:
    if not isinstance(value, list):
        raise ValueError(f"'{field_name}' must be an array of column names.")
    return tuple(_nonempty_string(item, field_name) for item in value)


def _nonempty_string(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"'{field_name}' must be a non-empty string.")
    return value.strip()


def _compare_columns(expected: list[str], actual: list[str]) -> dict[str, list[str]]:
    return {
        "missing": [column for column in expected if column not in actual],
        "extra": [column for column in actual if column not in expected],
    }


def _normalize_target_type(target_type: str) -> str:
    aliases = {
        "str": "string",
        "text": "string",
        "char": "string",
        "varchar": "string",
        "nvarchar": "string",
        "nchar": "string",
        "character varying": "string",
        "int": "integer",
        "bigint": "integer",
        "smallint": "integer",
        "float": "number",
        "double": "number",
        "double precision": "number",
        "real": "number",
        "numeric": "decimal",
        "bool": "boolean",
        "bit": "boolean",
        "datetime": "date",
        "timestamp": "date",
        "timestamp with time zone": "date",
        "timestamp without time zone": "date",
    }
    normalized = target_type.strip().lower()
    if normalized in aliases:
        return aliases[normalized]
    base_type = normalized.split("(", maxsplit=1)[0].strip()
    return aliases.get(base_type, base_type)


def _apply_operation(operation: str, values: list[Any], options: Mapping[str, Any]) -> Any:
    if operation == "copy":
        _require_value_count(operation, values, 1)
        return values[0]
    if operation in {"trim", "upper", "lower"}:
        _require_value_count(operation, values, 1)
        value = values[0]
        if value is None:
            return None
        if not isinstance(value, str):
            raise ValueError(f"Transformation '{operation}' requires a string value.")
        return {
            "trim": str.strip,
            "upper": str.upper,
            "lower": str.lower,
        }[operation](value)
    if operation == "concat":
        separator = options.get("separator", "")
        if not isinstance(separator, str):
            raise ValueError("Transformation 'concat' requires a string separator.")
        if any(value is None for value in values):
            return None
        return separator.join(str(value) for value in values)
    if operation == "coalesce":
        for value in values:
            if value is not None:
                return value
        return None
    if operation == "replace":
        _require_value_count(operation, values, 1)
        value = values[0]
        if value is None:
            return None
        old = options.get("old")
        new = options.get("new", "")
        if not isinstance(value, str) or not isinstance(old, str) or not isinstance(new, str):
            raise ValueError("Transformation 'replace' requires a string input and string old/new options.")
        return value.replace(old, new)
    if operation == "substring":
        _require_value_count(operation, values, 1)
        value = values[0]
        if value is None:
            return None
        if not isinstance(value, str):
            raise ValueError("Transformation 'substring' requires a string value.")
        start = options.get("start", 0)
        length = options.get("length")
        if not isinstance(start, int) or (length is not None and not isinstance(length, int)):
            raise ValueError("Transformation 'substring' requires integer start and length options.")
        return value[start:] if length is None else value[start:start + length]
    if operation == "round":
        _require_value_count(operation, values, 1)
        value = values[0]
        if value is None:
            return None
        digits = options.get("digits", 0)
        if not isinstance(digits, int):
            raise ValueError("Transformation 'round' requires an integer digits option.")
        try:
            return round(Decimal(str(value)), digits)
        except (InvalidOperation, TypeError, ValueError) as error:
            raise ValueError("Transformation 'round' requires a numeric value.") from error
    if operation == "if_else":
        _require_value_count(operation, values, 1)
        expected = options.get("equals")
        return options.get("then") if values[0] == expected else options.get("else")
    if operation in _ARITHMETIC_OPERATIONS:
        if operation == "subtract" or operation == "divide":
            _require_value_count(operation, values, 2)
            operands = values
        else:
            if len(values) < 2:
                raise ValueError(f"Transformation '{operation}' requires at least two source columns.")
            operands = values
        if any(value is None for value in operands):
            return None
        try:
            numbers = [Decimal(str(value)) for value in operands]
            result = numbers[0]
            if operation in {"subtract", "divide"}:
                return _ARITHMETIC_OPERATIONS[operation](numbers[0], numbers[1])
            for value in numbers[1:]:
                result = _ARITHMETIC_OPERATIONS[operation](result, value)
            return result
        except (InvalidOperation, TypeError, ZeroDivisionError) as error:
            raise ValueError(f"Transformation '{operation}' could not be applied to its inputs.") from error
    raise ValueError(f"Unsupported transformation '{operation}'.")


def _require_value_count(operation: str, values: list[Any], expected: int) -> None:
    if len(values) != expected:
        raise ValueError(
            f"Transformation '{operation}' requires exactly {expected} source column(s); got {len(values)}."
        )


def _cast_value(value: Any, target_type: str, target_column: str) -> Any:
    try:
        if target_type == "string":
            return str(value)
        if target_type == "integer":
            integer_value = Decimal(str(value))
            if integer_value != integer_value.to_integral_value():
                raise ValueError("value is not an integer")
            return int(integer_value)
        if target_type == "number":
            return float(value)
        if target_type == "decimal":
            return Decimal(str(value))
        if target_type == "boolean":
            if isinstance(value, bool):
                return value
            normalized = str(value).strip().lower()
            if normalized in {"true", "1", "yes"}:
                return True
            if normalized in {"false", "0", "no"}:
                return False
            raise ValueError("expected true/false, 1/0, or yes/no")
        if target_type == "date":
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            return date.fromisoformat(str(value))
    except (ValueError, TypeError, InvalidOperation) as error:
        raise ValueError(
            f"Value for target '{target_column}' cannot be converted to '{target_type}'."
        ) from error
    raise ValueError(f"Unsupported target type '{target_type}'.")
